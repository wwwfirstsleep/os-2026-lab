#!/usr/bin/env python3
"""Build the Agent-verified DOCX from current source and immutable/raw logs.
Requires python-docx. Export DOCX to PDF with Word/LibreOffice after rendering QA.
"""
from pathlib import Path
import re, textwrap
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
D=Document();S=D.sections[0]
S.page_width=Cm(21);S.page_height=Cm(29.7)
S.top_margin=S.bottom_margin=Cm(2.1);S.left_margin=S.right_margin=Cm(2.2)
F='Noto Sans SC'
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 st=D.styles[name];st.font.name=F;st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),F)
 st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.18
D.styles['Normal'].font.size=Pt(10)
for name,size in [('Title',24),('Heading 1',15),('Heading 2',12)]:
 D.styles[name].font.size=Pt(size);D.styles[name].paragraph_format.keep_with_next=True
for el in list(D.styles._element.xpath('.//w:pBdr')):el.getparent().remove(el)
f=S.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
def p(t):return D.add_paragraph(t)
def h(t):return D.add_heading(t,1)
def page():D.add_page_break()
def code(t,size=8.3,width=95):
 for line in t.splitlines():
  for wrap in textwrap.wrap(line,width,replace_whitespace=False,drop_whitespace=False,subsequent_indent='    ') or ['']:
   q=D.add_paragraph();q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=Pt(size+2)
   q.paragraph_format.keep_with_next=False
   r=q.add_run(wrap);r.font.name='DejaVu Sans Mono';r.font.size=Pt(size)
   r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),F)
def log_case(name,filename='final_test.log'):
 text=(R/'logs'/filename).read_text()
 match=re.search(r'^CASE '+re.escape(name)+r'\n.*?^PASS$',text,re.M|re.S)
 assert match,name
 return match.group(0)
for i,t in enumerate(['操作系统','实验一','Linux 命令解释程序设计与实现']):
 q=D.add_paragraph(t,'Title');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
 q.paragraph_format.space_before=Pt(45 if i==0 else 10)
for i,t in enumerate(['实验报告','学期 2026—2027第一学期','学院 网络空间安全学院','姓名 王旭升','学号 24281204','班级 保密2401','编制日期 2026年10月2日']):
 q=p(t);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_before=Pt(22 if i==0 else 3)
p('开发方式：Agent 在 Linux 中实际编写和验证；报告如实保留协助说明与功能边界。')
page()
sections=[];current=None
for line in (R/'docs/report.md').read_text().splitlines():
 if line.startswith('## '):current=[line[3:],[]];sections.append(current)
 elif current and line.strip():current[1].append(line)
for title,paras in sections:
 if title in ['实验步骤','实验结果','讨论与问题处理']:
  page()
 h(title)
 for text in paras:p(text)
 if title=='实验结果':
  p('完整测试和对应关系随项目提供。下方证据页摘录最终复验日志；长行只作排版折行。')
# Evidence pages use actual raw snippets, never manually invented output.
page();h('证据一 编译与最终综合测试')
p('来源：logs/final_build.log、logs/final_test.log、logs/verification_summary.log。')
code((R/'logs/final_build.log').read_text())
text=(R/'logs/final_test.log').read_text();code(text[text.index('PTY TRANSCRIPT:'):])
code((R/'logs/verification_summary.log').read_text())
p('69项是真实命令进程测试；PTY检查单列。4个解析器样例是合成日志输入，不冒充系统运行次数。')
page();h('证据二 内部命令与环境变量')
p('来源：logs/final_test.log。异常输出保持原样。')
for name in ['cd parent','HOME PATH expansion','cat file','rm error']:
 code(log_case(name));p('')
page();h('证据三 I/O 重定向与管道')
p('来源：logs/final_test.log。文件内容在测试脚本中另有断言。')
for name in ['output append input','echo write failure','pipe external','large pipe no deadlock','early reader exit']:
 code(log_case(name));p('')
page();h('证据四 实际系统调用')
p('来源：logs/strace_key.log、logs/strace_reaping.log。保留原PID、参数和返回值；完整上下文在 logs/strace.log。')
code((R/'logs/strace_key.log').read_text(),7.7)
code((R/'logs/strace_reaping.log').read_text(),7.7)
page();h('功能与测试证据对应表')
rows=[
 ('提示符与异常恢复','main','PTY TRANSCRIPT','final_test.log'),
 ('cd / pwd','builtin','cd parent / pwd','final_test.log；strace_cd_pwd.log'),
 ('echo','builtin','echo spaces / excessive arguments','basic_test.log'),
 ('cat','builtin','cat file / cat stdin / cat error','basic_test.log；strace_files.log'),
 ('touch / rm','builtin','touch create/preserve；rm file/error','basic_test.log；strace_files.log'),
 ('exit','builtin','exit arguments；pipe exit isolated','basic_test.log；pipe_test.log'),
 ('HOME / PATH','main / lex / execute','HOME PATH expansion；custom PATH','basic_test.log；strace_exec.log'),
 ('外部命令','execute','ls、whoami、uname；成功execve','basic_test.log；strace_exec.log'),
 ('> / >> / <','redirect / run_command','output append input；truncate','redirect_test.log；strace_key.log'),
 ('单管道','run_line / pipeline_child','pipe external；builtin；large pipe','pipe_test.log；strace_pipe.log'),
 ('fd关闭与回收','run_command / wait_child','100轮 fd probe；实际 wait4','final_test.log；strace_reaping.log'),
 ('解析边界与限制','lex / parse','line/token limit；unsupported glob','final_test.log'),
 ('摘录脚本修复','extract_trace.py','4个合成回归样例','trace_parser_test.log')]
t=D.add_table(rows=1,cols=4);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
widths=[3.1,3.5,5.3,4.6]
for i,w in enumerate(widths):t.columns[i].width=Cm(w)
for c,txt in zip(t.rows[0].cells,['功能','代码函数','测试名称或操作','证据文件（logs/）']):c.text=txt
for row in rows:
 for c,txt in zip(t.add_row().cells,row):c.text=txt
for row in t.rows:
 for i,c in enumerate(row.cells):
  c.width=Cm(widths[i]);c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
  for q in c.paragraphs:
   q.paragraph_format.space_after=Pt(5);q.paragraph_format.space_before=Pt(5)
   for run in q.runs:run.font.size=Pt(8.5)
  pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
  pr.append(b)
for c in t.rows[0].cells:
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAF0F5');c._tc.get_or_add_tcPr().append(sh)
 for q in c.paragraphs:
  for run in q.runs:run.bold=True
rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
p('所有C函数位于 src/myshell.c；完整命令与断言位于 tests/run_tests.py。')
page();h('关键原理说明')
for line in (R/'docs/understanding.md').read_text().splitlines():
 if line.strip() and not line.startswith('#'):p(re.sub(r'\*\*(.*?)\*\*',r'\1',line))
page();h('附录 源程序完整清单')
p('当前 src/myshell.c，共333行；左侧为源文件行号，长行仅视觉折行。')
code('\n'.join(f'{i:03d} {line}' for i,line in enumerate((R/'src/myshell.c').read_text().splitlines(),1)),7.7,99)
page();h('附录 Makefile 与复验入口')
code((R/'Makefile').read_text().replace('\t','    '),10)
p('实际Makefile配方使用Tab。复验入口：bash tests/verify.sh。')
code((R/'tests/verify.sh').read_text(),8.1)
h('当前编译对象校验值')
code((R/'logs/artifact_hashes.log').read_text(),7.7)
p('完整聊天导出与 GitHub 远程地址尚未取得；详见 docs/AGENT_DISCLOSURE.md 和 docs/GITHUB_SUBMISSION.md。本报告不将这些待办写成已完成。')
D.core_properties.title='Linux命令解释程序设计与实现实验报告'
D.core_properties.author='王旭升'
D.core_properties.comments='AI-assisted; actual Linux execution evidence; complete chat export pending.'
D.save(R/'docs/report.docx')
print(R/'docs/report.docx')
