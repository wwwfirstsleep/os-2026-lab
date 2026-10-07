from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
R=Path(__file__).resolve().parents[1]
D=Document(R/'docs/report.docx')
labels=[
('01-environment','图1 当前Linux环境与工具版本'),('02-build','图2 清理后重新编译成功'),
('03-prompt','图3 Shell启动及真实PTY提示符'),('04-builtins','图4 pwd、echo与HOME展开'),
('05-external','图5 外部命令运行'),('06-cd','图6 主进程目录改变与错误后继续'),
('07-files','图7 文件创建、读取、删除及异常恢复'),('08-redirection','图8 实际输出、追加和输入重定向'),
('09-pipe','图9 单级管道，printf hello计数为5'),('10-strace-history','图10 查看历史成功strace原始日志（非本次新跟踪）'),
('11-git','图11 截图时的真实Git开发历史'),('12-tests','图12 本次测试：68项后受/proc限制，追加18项通过'),
('13-strace-current','图13 本次strace权限拒绝，保留真实失败'),('14-test-history','图14 查看历史完整通过日志并核对当前编译对象哈希')]
for i,(file,label) in enumerate(labels):
 if i%2==0:
  D.add_page_break()
  D.add_heading('补充 实际终端截图' if i==0 else '实际终端截图 续',level=1)
  if i==0:
   q=D.add_paragraph('2026年10月7日UTC补充。以下为Agent在Linux中运行真实PTY与xterm后截取的窗口，含姓名拼音与学号。图10、14明确查看历史记录；图12、13保留本次环境限制，未声称本次完整复验成功。')
   for run in q.runs:run.font.size=Pt(9)
 q=D.add_paragraph(label);q.paragraph_format.space_before=Pt(2);q.paragraph_format.space_after=Pt(2)
 for run in q.runs:run.font.size=Pt(9);run.bold=True
 q.paragraph_format.keep_with_next=True
 q=D.add_paragraph();q.paragraph_format.space_after=Pt(3)
 q.add_run().add_picture(str(R/'docs/screenshots'/f'{file}.png'),width=Cm(16.1))
D.save(R/'docs/report-with-screenshots.docx')
print('Saved report-with-screenshots.docx, 14 images')
