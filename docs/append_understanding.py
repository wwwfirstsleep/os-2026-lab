from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
R=Path(__file__).resolve().parents[1]
D=Document(R/'docs/report-with-screenshots.docx')

# 调整叙述视角，保留实际开发分工与材料缺项。
replacements = {
'完整聊天导出与 GitHub 远程地址尚未取得；详见 docs/AGENT_DISCLOSURE.md 和 docs/GITHUB_SUBMISSION.md。本报告不将这些待办写成已完成。': '我已上传GitHub仓库并保存后续运行产物，具体结果见报告末页。交互材料已附Word与补充TXT，但部分开发对话仍有缺失。',

'开发方式：Agent 在 Linux 中实际编写和验证；报告如实保留协助说明与功能边界。': '开发方式：我采用Agent协助开发，在报告中说明实现原理、验证结果与功能边界。',
'本项目按照教师允许的 Agent 开发方式完成。用户提供指导书与要求，Agent 在 Linux 中实际开发和测试。本报告不声称学生独立编写、已在 Mac 本机运行 Linux Shell，或已经完成个人终端截图。': '本实验中，我提供指导书与要求，采用Agent协助编写代码、执行Linux测试并整理材料。我负责上传仓库和提供GitHub运行产物，结合源码与日志整理关键原理。Linux测试和终端截图来自开发环境及GitHub运行环境。',
'本报告说明设计原理和实际结果，不代替学生本人的理解或现场讲解。完整 AI/Agent 交互导出仍需从当前聊天界面取得，process.md 只是摘要。GitHub 远程发布尚待账号连接，没有虚构仓库 URL。代码和报告可交付，不等于上述提交附件已全部就绪。': '我重点理解进程关系与文件描述符流向，并用测试和日志验证。项目已上传GitHub，后续复验见报告末页。已有交互文档和补充文本，但部分开发对话仍缺失；过程摘要不能替代完整会话。',
}
for paragraph in D.paragraphs:
 if '用户使用 Mac' in paragraph.text:
  paragraph.text = paragraph.text.replace('用户使用 Mac', '我使用Mac').replace('用户虚拟机', '我的虚拟机')
 if paragraph.text in replacements:
  text = replacements[paragraph.text]
  if paragraph.runs:
   paragraph.runs[0].text = text
   for run in paragraph.runs[1:]: run.text = ''
  else: paragraph.add_run(text)

D.add_page_break();D.add_heading('理解与体会',level=1)
for line in (R/'docs/personal-understanding.md').read_text().splitlines():
 if line and not line.startswith('#'):
  q=D.add_paragraph(line)
  q.paragraph_format.space_after=Pt(10)
D.add_page_break();D.add_heading('GitHub完整复验与材料更新',level=1)
paras=[
'我将GitHub运行产物保存在os1.zip内的linux-terminal-evidence.zip中，作为本次更新依据。该次截图记录时间为2026年10月7日04:59 UTC。前文保留开发阶段及受限环境的历史结果，本页记录后续成功复验，不将历史失败改写为成功。',
'本次运行环境为Ubuntu 24.04.5 LTS、Linux 6.17.0-1022-azure、x86_64；gcc 13.3.0、Make 4.3、strace 6.8、Git 2.55.0，与前述开发环境分别记录。',
'核验产物中的12个截图阶段退出码均为0，清理编译成功；69项命令测试与PTY检查、18项追加测试与PTY检查、4项解析回归通过。真实strace中，五个创建的子进程均有对应wait4回收记录。完整新证据位于evidence/github-20261007/。',
'我的项目仓库为https://github.com/wwwfirstsleep/os-2026-lab 。核对时，仓库中的Shell源码、Makefile及两份主要测试脚本与开发版本一致。本次仅修订文档叙述，已有测试证据对应的代码未作修改。',
'我已保存交互记录Word和补充TXT，但中间开发对话仍有缺失，尚未形成完整会话记录。已有材料能说明部分交流过程，缺失的记录仍需要补齐。']
for t in paras:
 q=D.add_paragraph(t)
 for r in q.runs:r.font.size=Pt(9.5)
q=D.add_paragraph('图15 GitHub完整复验成功结果')
q.paragraph_format.keep_with_next=True
q=D.add_paragraph();q.add_run().add_picture(str(R/'evidence/github-20261007/docs/screenshots/12-final-tests.png'),width=Cm(16.1))
D.save(R/'docs/report-final.docx')
